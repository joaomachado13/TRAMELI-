"""Prepare a private SQL import from the supplier tab. Never commit its output."""

import argparse
import json
import unicodedata
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


def normalized(value):
    plain = unicodedata.normalize("NFD", str(value).strip().lower())
    return " ".join("".join(char for char in plain if unicodedata.category(char) != "Mn").split())


def quote(value):
    return "'" + str(value).replace("'", "''") + "'"


def cents(value):
    amount = Decimal(str(value)) * 100
    if amount != amount.to_integral_value() or amount < 0 or amount > 100000000:
        raise ValueError(f"Preço da padaria inválido: {value!r}")
    return int(amount)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="DADOS PADARIA.xlsx")
    parser.add_argument("--output", type=Path, default=Path("private/supplier-cost-import.sql"))
    args = parser.parse_args()

    catalog = json.loads(Path("data/client-products.json").read_text(encoding="utf-8"))["products"]
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    if "Valores Padaria" not in workbook.sheetnames:
        raise ValueError("A aba 'Valores Padaria' não foi encontrada.")
    supplier = defaultdict(list)
    for row_number, row in enumerate(workbook["Valores Padaria"].iter_rows(values_only=True), 1):
        if row_number == 1 or not isinstance(row[0], str) or not isinstance(row[1], (int, float)):
            continue
        name = row[0].strip()
        supplier[normalized(name)].append((row_number, name, cents(row[1])))

    matches = []
    for product in catalog:
        found = supplier.get(normalized(product["name"]), [])
        if len(found) == 1:
            source_row, supplier_name, cost = found[0]
            matches.append((product["id"], cost, supplier_name, source_row))

    if not matches:
        raise ValueError("Nenhuma correspondência única; verifique a planilha e o catálogo.")
    values = [f"  ({quote(product_id)}, {cost}, {quote(name)}, {row})"
              for product_id, cost, name, row in matches]
    sql = (
        "-- PRIVADO: gerado de DADOS PADARIA.xlsx / Valores Padaria.\n"
        "-- Não adicionar este arquivo ao Git. Revisar correspondências antes de aplicar.\n"
        f"-- {len(matches)} correspondências únicas por nome normalizado; "
        f"{len(catalog) - len(matches)} sem custo confirmado.\n"
        "begin;\n"
        "insert into public.trameli_product_costs\n"
        "  (product_id, unit_cost_cents, supplier_name, supplier_source_row)\n"
        "values\n" + ",\n".join(values) + "\n"
        "on conflict (product_id) do update set\n"
        "  unit_cost_cents = excluded.unit_cost_cents,\n"
        "  supplier_name = excluded.supplier_name,\n"
        "  supplier_source_row = excluded.supplier_source_row,\n"
        "  updated_at = now();\n"
        "commit;\n"
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(sql, encoding="utf-8")
    print(f"Arquivo privado preparado: {len(matches)} custos associados; "
          f"{len(catalog) - len(matches)} pendentes. Revise antes de aplicar.")


if __name__ == "__main__":
    main()
